# -*- coding: utf-8 -*-
"""
Full Diploma Thesis Builder for "Незбіг"
Topic: Програмна система інтелектуального аналізу текстових документів для виявлення запозичень та штучно згенерованого контенту
Specialty: 121 Інженерія програмного забезпечення
Output: docs/ДИПЛОМНА_РОБОТА_НЕЗБІГ.docx
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_run_font(run, font_name="Times New Roman", size_pt=14, bold=False, italic=False, color_rgb=(0, 0, 0)):
    """
    Strictly forces font family, size and scripts in Word OOXML
    ensuring Times New Roman is applied for Latin, Cyrillic, numbers, and symbols.
    """
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    
    half_points = str(int(size_pt * 2))
    sz = rPr.find(qn('w:sz'))
    if sz is None:
        sz = OxmlElement('w:sz')
        rPr.append(sz)
    sz.set(qn('w:val'), half_points)
    
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = OxmlElement('w:szCs')
        rPr.append(szCs)
    szCs.set(qn('w:val'), half_points)

def create_document():
    doc = Document()
    
    # Page setup (A4, standard Ukrainian academic margins)
    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(3.0)
        s.right_margin = Cm(1.0)
        s.different_first_page_header_footer = True
        
        # Header setup: page number top right
        header = s.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.paragraph_format.space_before = Pt(0)
        p_head.paragraph_format.space_after = Pt(0)
        
        # Add dynamic page number field in OOXML
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:cs'), 'Times New Roman')
        rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '28') # 14 pt
        szCs = OxmlElement('w:szCs')
        szCs.set(qn('w:val'), '28')
        rPr.append(rFonts)
        rPr.append(sz)
        rPr.append(szCs)
        r.append(rPr)
        fldSimple.append(r)
        p_head._p.append(fldSimple)

    # Style: Normal
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(14)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(0)
    style_normal.paragraph_format.space_before = Pt(0)
    style_normal.paragraph_format.first_line_indent = Cm(1.25)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    return doc

def add_p(doc, text="", bold_prefix=None, space_after=0, space_before=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=1.25, line_spacing=1.5):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if first_indent is not None:
        p.paragraph_format.first_line_indent = Cm(first_indent)
    else:
        p.paragraph_format.first_line_indent = Cm(0)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, font_name="Times New Roman", size_pt=14, bold=True)
        
    if text:
        r = p.add_run(text)
        set_run_font(r, font_name="Times New Roman", size_pt=14, bold=False)
    return p

def add_chapter_heading(doc, text, new_page=True):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    
    r = p.add_run(text.upper())
    set_run_font(r, font_name="Times New Roman", size_pt=14, bold=True)
    return p

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.keep_with_next = True
    
    r = p.add_run(text)
    set_run_font(r, font_name="Times New Roman", size_pt=14, bold=True)
    return p

def add_subsection_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.keep_with_next = True
    
    r = p.add_run(text)
    set_run_font(r, font_name="Times New Roman", size_pt=14, bold=True, italic=True)
    return p

def add_formula(doc, formula_text, formula_num):
    """
    Renders academic formula using a clean borderless table:
    formula centered in Times New Roman 14 pt italic,
    formula number (num) right-aligned in Times New Roman 14 pt.
    """
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Remove borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)
    
    row = table.rows[0]
    row.cells[0].width = Cm(14.5)
    row.cells[1].width = Cm(2.5)
    
    # Formula cell (center)
    p_f = row.cells[0].paragraphs[0]
    p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f.paragraph_format.line_spacing = 1.5
    p_f.paragraph_format.space_before = Pt(4)
    p_f.paragraph_format.space_after = Pt(4)
    p_f.paragraph_format.first_line_indent = Cm(0)
    
    r1 = p_f.add_run(formula_text)
    set_run_font(r1, font_name="Times New Roman", size_pt=14, italic=True)
    
    # Number cell (right)
    p_n = row.cells[1].paragraphs[0]
    p_n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_n.paragraph_format.line_spacing = 1.5
    p_n.paragraph_format.space_before = Pt(4)
    p_n.paragraph_format.space_after = Pt(4)
    p_n.paragraph_format.first_line_indent = Cm(0)
    
    r2 = p_n.add_run(f"({formula_num})")
    set_run_font(r2, font_name="Times New Roman", size_pt=14, bold=False, italic=False)
    
    # Subtle trailing spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)
    p_after.paragraph_format.line_spacing = 1.0
    p_after.paragraph_format.first_line_indent = Cm(0)
    return table

def add_figure(doc, img_path, caption, width_cm=15.5):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found!")
        return None
    p_img = doc.add_paragraph()
    p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.first_line_indent = Cm(0)
    p_img.paragraph_format.keep_with_next = True
    
    r_img = p_img.add_run()
    r_img.add_picture(img_path, width=Cm(width_cm))
    
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.first_line_indent = Cm(0)
    
    r_cap = p_cap.add_run(caption)
    set_run_font(r_cap, font_name="Times New Roman", size_pt=14, bold=False, italic=False)
    return p_cap

def add_table_custom(doc, headers, data, caption=None, col_widths=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.first_line_indent = Cm(1.25)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(caption)
        set_run_font(r_cap, font_name="Times New Roman", size_pt=14, bold=True)
    
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)
    
    # Headers
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        for r in p.runs:
            set_run_font(r, font_name="Times New Roman", size_pt=13, bold=True)
        
        # Header background
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>')
        tcPr.append(shd)
        
    # Data rows
    for r_idx, row in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 or len(str(val)) > 15 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.first_line_indent = Cm(0)
            for r in p.runs:
                set_run_font(r, font_name="Times New Roman", size_pt=13, bold=False)
                
    # Column widths
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Cm(w)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)
    p_after.paragraph_format.first_line_indent = Cm(0)
    p_after.paragraph_format.line_spacing = 1.0

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.first_line_indent = Cm(1.25)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(caption)
        set_run_font(r_cap, font_name="Times New Roman", size_pt=14, bold=False, italic=True)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.rows[0].cells[0].width = Cm(16.5)
    
    cell = tbl.rows[0].cells[0]
    cell.text = code_str.strip()
    
    # Border & shading
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="008080"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8F9FA"/>')
    tcPr.append(shd)
    
    # Style paragraphs in code
    for p in cell.paragraphs:
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)
    p_after.paragraph_format.first_line_indent = Cm(0)

print("Formatting helpers ready.")
