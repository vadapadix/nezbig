# -*- coding: utf-8 -*-
"""
Thesis Chapter 3 Module: Software Implementation of the "Незбіг" System
All formatted strictly in Times New Roman 14 pt with proper academic layout.
Written in authentic, direct Ukrainian engineering prose without AI clichés or slop.
"""

import os
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_full_diploma_docx import (
    add_chapter_heading, add_section_heading, add_subsection_heading, 
    add_p, add_table_custom, add_figure, add_code_block
)

def add_chapter_3(doc):
    add_chapter_heading(doc, "РОЗДІЛ 3. ПРОГРАМНА РЕАЛІЗАЦІЯ СИСТЕМИ", new_page=True)
    
    # 3.1
    add_section_heading(doc, "3.1. Обґрунтування технологічного стеку та середовища розробки")
    
    add_p(doc, 
        text="Програмна реалізація вебсистеми «Незбіг» базується на сучасному стеку технологій, що забезпечує "
             "високу швидкодію, наскрізну типізацію, реактивність інтерфейсу та зручність супроводу коду. "
             "Вся кодова база проєкту реалізована мовою TypeScript (версія 5.9.3), що дозволяє використовувати "
             "спільні схеми типів як на серверному, так і на клієнтському рівні (`src/shared/types.ts`).")
             
    add_p(doc, 
        text="Клієнтський рівень побудований на базі бібліотеки React 19.2.3. Застосування найновішої версії React "
             "забезпечує ефективне використання Concurrent Mode та оптимізованого хука `useTransition` для опрацювання "
             "великих обсягів тексту без блокування основного потоку браузера. Збірка клієнтського застосунку здійснюється "
             "за допомогою Vite 7.3, що забезпечує миттєве гаряче оновлення модулів (HMR) під час розробки та створює "
             "компактні оптимізовані бандли для продакшн-середовища. Стилізація інтерфейсу реалізована через утилітарний "
             "фреймворк TailwindCSS 3.4 із кастомною темною темою оформлення та адаптивною сіткою.")
             
    add_p(doc, 
        text="Серверна частина функціонує на платформі Node.js версії 24 з використанням вебфреймворку Express 5.2. "
             "П'ята версія Express підтримує нативну обробку асинхронних маршрутів через проміси, що спрощує централізовану "
             "обробку винятків. Для аналізу документів використано бібліотеки Mammoth 1.11 (екстракція вмісту DOCX у санітизований HTML), "
             "pdf-parse 2.4 (вилучення тексту з файлів PDF) та JSZip 3.10 (низькорівнева маніпуляція пакетами Office Open XML). "
             "Парсинг та очищення HTML-коду завантажених вебсторінок виконується парсером Cheerio 1.1. "
             "Безпека забезпечується пакетами Helmet 8.3, DOMPurify 3.4, CORS та express-rate-limit.")

    headers_3_1 = ["Рівень системи", "Технологія / Бібліотека", "Версія", "Функціональне призначення"]
    data_3_1 = [
        ["Клієнт (Frontend)", "React", "19.2.3", "Побудова реактивного користувацького інтерфейсу (SPA)"],
        ["Клієнт (Frontend)", "TypeScript", "5.9.3", "Статична типізація, спільні інтерфейси та безпека даних"],
        ["Клієнт (Frontend)", "Vite", "7.3.0", "Високошвидкісний бандлер та інструмент збірки"],
        ["Клієнт (Frontend)", "TailwindCSS", "3.4.19", "Адаптивна верстка та людино-орієнтований темний UI"],
        ["Сервер (Backend)", "Node.js / Express", "24 / 5.2.1", "Асинхронний REST API сервер та маршрутизація"],
        ["Парсинг документів", "Mammoth / JSZip", "1.11 / 3.10", "Екстракція та реконструкція OOXML DOCX"],
        ["Парсинг документів", "pdf-parse", "2.4.5", "Вилучення тексту з вхідних файлів PDF"],
        ["Парсинг вебу", "Cheerio", "1.1.2", "Очищення HTML та виділення смислових блоків сторінок"],
        ["Пошукові провайдери", "Tavily / Serper API", "REST v1", "Багатопровайдерний вебіндекс без блокувань та капч"],
        ["Наукові бази", "Semantic Scholar / OpenAlex", "REST v1/v2", "Пошук за відкритими академічними каталогами та DOI"],
        ["Зовнішні LLM", "OpenRouter / NIM API", "REST v1", "Асинхронне отримання експертної оцінки (Llama 3.3)"],
        ["Тестування", "Vitest / Playwright", "4.0 / 1.62", "Модульне, інтеграційне та E2E тестування"]
    ]
    add_table_custom(doc, headers_3_1, data_3_1, 
                     caption="Таблиця 3.1 — Технологічний стек програмної системи «Незбіг»", 
                     col_widths=[3.5, 3.8, 2.2, 7.0])

    # 3.2
    add_section_heading(doc, "3.2. Архітектура сервера та проєктування REST API специфікації")
    
    add_p(doc, 
        text="Серверний бекенд побудований за принципом слабозв'язаних модулів. Усі кінцеві точки реалізують строгу "
             "валідацію вхідних даних за допомогою бібліотеки Zod 4.4, обмежують розмір тіла запиту (до 50 МБ) та повертають "
             "типізовані JSON-відповіді з діагностичними метриками.")

    headers_3_2 = ["HTTP Метод", "Маршрут API", "Формат тіла", "Опис функціоналу та відповідь"]
    data_3_2 = [
        ["POST", "/api/scan/jobs", "JSON {text, fileName, settings}", "Створення фонового асинхронного завдання сканування; повертає jobId"],
        ["GET", "/api/scan-status/:id", "URL Parameter id", "Опитування прогресу та отримання підсумкового звіту ScanReport"],
        ["POST", "/api/scan-file/jobs", "Multipart (file, settings)", "Завантаження DOCX/PDF та ініціалізація перевірки в пам'яті"],
        ["POST", "/api/humanize", "JSON {text, html, mode}", "Стилістичне редагування тексту (режими: academic, natural, concise)"],
        ["POST", "/api/humanize-file", "Multipart (file, mode)", "Стилістичне редагування Word із повним збереженням OOXML-структури"],
        ["POST", "/api/ai-opinion", "JSON {text, reportId}", "Асинхронний запит експертної думки зовнішньої моделі Meta Llama 3.3"],
        ["GET", "/api/health", "None", "Перевірка працездатності сервера та стану пошукових провайдерів"]
    ]
    add_table_custom(doc, headers_3_2, data_3_2, 
                     caption="Таблиця 3.2 — Специфікація REST API маршрутів системи «Незбіг»", 
                     col_widths=[2.3, 4.2, 3.6, 6.4])

    # 3.3
    add_section_heading(doc, "3.3. Програмна реалізація модулів аналізу та алгоритмів")
    
    add_subsection_heading(doc, "3.3.1. Реалізація алгоритмів скорингу та Winnowing")
    add_p(doc, 
        text="Модуль `src/server/winnowing.ts` реалізує документно-орієнтований фінгерпринтинг на основі поліноміального "
             "хешу Рабіна-Карпа з циклічним оновленням за O(1). Віконний фільтр обирає крайній правий мінімум, мінімізуючи обсяг "
             "збережених відбитків без втрати чутливості до локальних запозичень.")

    code_winnowing = """// src/server/winnowing.ts - Реалізація алгоритму Winnowing
export function winnowing(tokens: string[], k = 5, w = 4): Fingerprint[] {
  if (tokens.length < k) return [];
  const hashes: number[] = [];
  const B = 31;
  const M = 2147483647; // 2^31 - 1
  let power = 1;
  for (let i = 0; i < k - 1; i++) power = (power * B) % M;

  // 1. Поліноміальне хешування k-грам методом Рабіна-Карпа
  let currentHash = 0;
  for (let i = 0; i < k; i++) {
    currentHash = (currentHash * B + hashToken(tokens[i])) % M;
  }
  hashes.push(currentHash);

  for (let i = 1; i <= tokens.length - k; i++) {
    const prev = hashToken(tokens[i - 1]);
    const next = hashToken(tokens[i + k - 1]);
    currentHash = (currentHash - ((prev * power) % M) + M) % M;
    currentHash = (currentHash * B + next) % M;
    hashes.push(currentHash);
  }

  // 2. Ковзне вікно шириною w та вибір крайнього правого мінімуму
  const fingerprints: Fingerprint[] = [];
  let minIdx = -1;
  for (let i = 0; i <= hashes.length - w; i++) {
    let currentMin = hashes[i];
    let currentMinPos = i;
    for (let j = 1; j < w; j++) {
      if (hashes[i + j] <= currentMin) { // <= забезпечує rightmost minimum
        currentMin = hashes[i + j];
        currentMinPos = i + j;
      }
    }
    if (currentMinPos !== minIdx) {
      fingerprints.push({ hash: currentMin, position: currentMinPos });
      minIdx = currentMinPos;
    }
  }
  return fingerprints;
}"""
    add_code_block(doc, code_winnowing, caption="Лістинг 3.1 — Реалізація алгоритму Winnowing-фінгерпринтингу (TypeScript)")

    add_subsection_heading(doc, "3.3.2. Реалізація стилометричного AI-аналізатора та емпіричних маркерів")
    add_p(doc, 
        text="Модуль `src/server/aiStylometry.ts` здійснює розрахунок показників MATTR, Burstiness CV, а також перевірку "
             "трискладових переліків та діалектичного хеджування:")

    code_stylometry = """// src/server/aiStylometry.ts - Розрахунок стилометрії та емпіричних маркерів
export function computeStylometrics(tokens: string[], sentences: string[]): StylometricResult {
  // 1. Розрахунок ковзного лексичного багатства MATTR (вікно 50 слів)
  const windowSize = 50;
  let mattrScore = 1.0;
  if (tokens.length >= windowSize) {
    let totalTTR = 0;
    const steps = tokens.length - windowSize + 1;
    for (let i = 0; i < steps; i++) {
      const windowTokens = tokens.slice(i, i + windowSize);
      const uniqueTypes = new Set(windowTokens).size;
      totalTTR += uniqueTypes / windowSize;
    }
    mattrScore = totalTTR / steps;
  }

  // 2. Варіація темпоритму довжини речень (Burstiness CV)
  const lengths = sentences.map(s => s.trim().split(/\\s+/).filter(Boolean).length).filter(l => l > 0);
  const mean = lengths.reduce((acc, v) => acc + v, 0) / (lengths.length || 1);
  const variance = lengths.reduce((acc, v) => acc + Math.pow(v - mean, 2), 0) / (lengths.length || 1);
  const cv = mean > 0 ? Math.sqrt(variance) / mean : 0.5;

  // 3. Виявлення трискладових переліків (Tricolon) та діалектичного хеджування
  const tricolonPattern = /([\\wа-яіїєґ]+),\\s+([\\wа-яіїєґ]+)\\s+(?:та|і|й)\\s+([\\wа-яіїєґ]+)/gi;
  const hedgingPattern = /з одного боку|з іншого боку|разом з тим|варто зауважити/gi;
  const fullText = sentences.join(' ');
  const tricolonMatches = (fullText.match(tricolonPattern) || []).length;
  const hedgingMatches = (fullText.match(hedgingPattern) || []).length;

  return { mattr: mattrScore, burstinessCV: cv, tricolons: tricolonMatches, hedgings: hedgingMatches };
}"""
    add_code_block(doc, code_stylometry, caption="Лістинг 3.2 — Реалізація розрахунку MATTR, Burstiness CV та патернів LLM (TypeScript)")

    add_subsection_heading(doc, "3.3.3. Реалізація фільтрації коду та структури документів")
    add_p(doc, 
        text="Модуль `src/server/proseFilter.ts` здійснює розпізнавання кодових конструкцій. Рядки аналізуються на наявність "
             "ключових слів програмування (`function`, `class`, `import`, `return`, `SELECT`, `FROM`), характерних символів "
             "(фігурні дужки, оператори присвоєння) та характерних відступів. Кодові рядки відсікаються від стилометрії "
             "та обліковуються в `aiExclusions.codeWords`. Бібліографічний парсер відсікає списки літератури (`referenceWords`), "
             "а парсер титульних блоків усуває службові реквізити (`skippedTitleWords`).")

    add_subsection_heading(doc, "3.3.4. Реалізація маніпулятора документами Microsoft Word OOXML")
    add_p(doc, 
        text="Модуль `src/server/formattedDocx.ts` розпаковує ZIP-пакет DOCX за допомогою бібліотеки JSZip. "
             "Він парсить файл `word/document.xml`, перетворює його на DOM-дерево та викликає `textAlignment.ts`. "
             "Вирівнювач токенів зіставляє початкові слова зі зміненими після стилістичної обробки. "
             "Зміни вносяться виключно у вузли `<w:t>`, не зачіпаючи параметрів форматування `<w:rPr>` (шрифти, накреслення, "
             "кегль, колір). Сформований пакет упаковується назад у бінарний буфер DOCX зі збереженням вихідного макета.")

    # 3.4
    add_section_heading(doc, "3.4. Реалізація детермінованого стилістичного редактора (Text Humanizer)")
    
    add_p(doc, 
        text="Модуль `src/server/humanizer.ts` забезпечує цільове переписування тексту для усунення характерних ознак генеративного стилю. "
             "Він підтримує три режими опрацювання:")
    add_p(doc,
        text="1. `academic` — трансформує пасивні дієслівні форми на суб'єктний активний стан, усуває шаблонні вступні конструкції "
             "(«слід зазначити», «дослідження показують», «можна стверджувати»), оптимізує спеціальну термінологію;")
    add_p(doc,
        text="2. `natural` — усуває монотонність синтаксичного ритму шляхом чергування коротких тез і розгорнутих конструкцій (нормалізація Burstiness);")
    add_p(doc,
        text="3. `concise` — стискає тавтологічні повтори та канцеляризми, підвищуючи інформаційну щільність тексту.")
              
    add_p(doc, 
        text="Результат повертається клієнту у вигляді об'єкта `HumanizeResult`, що містить скоригований текст, показники "
             "ризику ШІ до і після редагування (наприклад, зменшення з 88% до 14%) та деталізований масив груп змін за категоріями: "
             "`cliche` (анти-кліше), `pacing` (темпоритм), `syntax` (синтаксис), `vocabulary` (лексика).")

    # 3.5
    add_section_heading(doc, "3.5. Клієнтський користувацький інтерфейс на React 19")
    
    add_p(doc, 
        text="Користувацький інтерфейс системи «Незбіг» спроєктовано за принципами людино-орієнтованого дизайну (Human-Centered Design) "
             "з нативною підтримкою української локалізації. Ключові компоненти клієнтського застосунку:")
    add_p(doc,
        text="1. `TextEditor.tsx` — робоча область введення тексту з лічильником слів та зоною Drag-and-Drop для завантаження файлів DOCX/PDF;")
    add_p(doc,
        text="2. `ScanSettingsPanel.tsx` — бічна панель налаштування глибини перевірки («Швидко», «Глибоко», «Експертно») "
             "із динамічною індикацією розміру фрагмента та оверлапу;")
    add_p(doc,
        text="3. `ReportView.tsx` — зведений звіт перевірки з Executive Hero Dashboard, вердикт-бейджем, числовими картками "
             "із смугою похибки (±N п.п.) та трьома сегментованими вкладками: «Джерела та збіги», «Аналіз ШІ та AI-думка», «Додаткові відомості»; "
             "містить кнопку повтору запиту до моделі та швидке повернення до редактора;")
    add_p(doc,
        text="4. `HumanizePanel.tsx` — панель перегляду результатів адаптивної стилізації з порівнянням ризику ШІ до і після обробки;")
    add_p(doc,
        text="5. `HistoryPage.tsx` — журнал збережених перевірок у локальному сховищі браузера з можливістю відкриття будь-якого звіту.")

    sc1_path = os.path.abspath("docs/screenshots/01_main_page_ua.png")
    add_figure(doc, sc1_path, "Рис. 3.1 — Головна робоча область системи «Незбіг»: завантаження документів та редактор", width_cm=15.5)

    sc2_path = os.path.abspath("docs/screenshots/02_settings_ua.png")
    add_figure(doc, sc2_path, "Рис. 3.2 — Панель параметрів сканування та вибору глибини перевірки", width_cm=15.5)

    sc3_path = os.path.abspath("docs/screenshots/03_humanizer_page.png")
    add_figure(doc, sc3_path, "Рис. 3.3 — Інтерфейс модуля адаптивного стилістичного редагування тексту (Text Humanizer)", width_cm=15.5)

    sc4_path = os.path.abspath("docs/screenshots/04_humanizer_diff_result.png")
    add_figure(doc, sc4_path, "Рис. 3.4 — Візуалізація результатів стилістичного редагування та зниження ризику ШІ", width_cm=15.5)

    sc5_path = os.path.abspath("docs/screenshots/05_history_page.png")
    add_figure(doc, sc5_path, "Рис. 3.5 — Панель журналу історії перевірок документів", width_cm=15.5)

    # 3.6
    add_section_heading(doc, "3.6. Висновки до розділу 3")
    
    add_p(doc, 
        text="1. Здійснено програмну реалізацію системи «Незбіг» на базі сучасного технологічного стеку React 19, TypeScript 5.9, "
             "Vite 7.3, Node.js 24 та Express 5.2 з єдиною системою типізації даних.")
    add_p(doc,
        text="2. Розроблено та задокументовано специфікацію REST API, що підтримує асинхронні завдання перевірки, обробку файлів DOCX/PDF "
             "та фонове опитування зовнішніх мовних моделей.")
    add_p(doc,
        text="3. Програмно реалізовано модулі Winnowing-фінгерпринтингу, 5-факторного скорингу та розрахунку стилометричних маркерів "
             "(MATTR, Burstiness CV, трискладові переліки, діалектичне хеджування).")
    add_p(doc,
        text="4. Створено модуль `formattedDocx.ts` на основі JSZip та токенного вирівнювання, що забезпечує 100% збереження структури OOXML "
             "при стилістичній правці файлів Microsoft Word.")
    add_p(doc,
        text="5. Реалізовано модуль адаптивної стилізації Text Humanizer з трьома режимами (Academic, Natural, Concise) та оновлений "
             "компонент ReportView з Executive Hero Dashboard і сегментованими вкладками аналітики.")

print("Chapter 3 module ready.")
